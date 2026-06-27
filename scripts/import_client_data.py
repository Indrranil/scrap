#!/usr/bin/env python3
"""Import client item and vendor data from Excel files.

Usage:
    python scripts/import_client_data.py \\
        --item-code upated_itemcode.xlsx \\
        --client-data updated_8digit.xlsx
"""

from __future__ import annotations

import argparse
import os
import sys
import time
from decimal import Decimal
from typing import Dict, List, Optional, Tuple

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from sqlalchemy.orm import Session

from app.constants.items import (
    DEFAULT_CLIENT_DATA_XLSX,
    DEFAULT_ITEM_CODE_XLSX,
    P_ITEM_SHRED_OUTPUTS,
    PLU_ITEM_CODE_OVERRIDES,
    SHREDDABLE_PLU_CODES,
    SKIP_VENDOR_ITEM_NAMES,
)
from app.database.connection import SessionLocal
from app.models.item_master import ItemMaster
from app.models.vendor import Vendor
from app.models.vendor_item import VendorItem
from app.services.item_master import normalize_plu_code

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit(
        "openpyxl is required. Install with: pip install openpyxl"
    ) from exc


def _normalize_name(name: str) -> str:
    return " ".join(name.upper().split())


def _load_plu_rows(path: str) -> List[Tuple[str, str, str]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["PLU"]
    rows: List[Tuple[str, str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[1]:
            continue
        plu = normalize_plu_code(
            str(int(row[1])) if isinstance(row[1], float) else str(row[1])
        )
        name = str(row[2]).strip()
        uom = str(row[3]).strip().upper()
        rows.append((plu, name, uom))
    wb.close()
    return rows


def _load_rates_pdf(path: str) -> Dict[str, str]:
    """Map normalized item name -> 8-digit code (exact names only)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Rates PDF"]
    mapping: Dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) < 2 or not row[0] or not row[1]:
            continue
        code = str(int(row[0]))
        name = str(row[1]).strip()
        mapping[_normalize_name(name)] = code
    wb.close()
    return mapping


def _resolve_item_code(
    plu_code: str,
    name: str,
    rates: Dict[str, str],
    assigned_codes: Dict[str, str],
) -> Tuple[Optional[str], Optional[str], Optional[str]]:
    """Return (item_code, shred_output_code, shred_output_name)."""
    if plu_code in P_ITEM_SHRED_OUTPUTS:
        out_code, out_name = P_ITEM_SHRED_OUTPUTS[plu_code]
        return None, out_code, out_name

    if plu_code in PLU_ITEM_CODE_OVERRIDES:
        code = PLU_ITEM_CODE_OVERRIDES[plu_code]
    else:
        code = rates.get(_normalize_name(name))

    if not code:
        return None, None, None

    owner = assigned_codes.get(code)
    if owner and owner != plu_code:
        print(
            f"  WARNING: PLU {plu_code} ({name}) skipped item_code {code} "
            f"— already used by PLU {owner}"
        )
        return None, None, None

    assigned_codes[code] = plu_code
    return code, None, None


def import_items(db: Session, item_code_path: str, client_data_path: str) -> None:
    plu_rows = _load_plu_rows(item_code_path)
    rates = _load_rates_pdf(client_data_path)
    now = int(time.time())
    unmatched: List[str] = []
    assigned_codes: Dict[str, str] = {}

    for plu_code, name, uom in plu_rows:
        item_code, shred_out_code, shred_out_name = _resolve_item_code(
            plu_code, name, rates, assigned_codes
        )
        is_p_item = plu_code in P_ITEM_SHRED_OUTPUTS
        if not is_p_item and not item_code:
            unmatched.append(f"{plu_code} | {name}")

        existing = db.query(ItemMaster).filter(ItemMaster.plu_code == plu_code).first()
        if existing:
            existing.name = name
            existing.uom = uom
            existing.item_code = item_code
            existing.is_p_item = is_p_item
            existing.shred_output_item_code = shred_out_code
            existing.shred_output_name = shred_out_name
            existing.is_shreddable = plu_code in SHREDDABLE_PLU_CODES
            existing.is_active = True
            existing.updated_at = now
        else:
            db.add(
                ItemMaster(
                    plu_code=plu_code,
                    item_code=item_code,
                    name=name,
                    uom=uom,
                    is_shreddable=plu_code in SHREDDABLE_PLU_CODES,
                    is_p_item=is_p_item,
                    shred_output_item_code=shred_out_code,
                    shred_output_name=shred_out_name,
                    is_active=True,
                    created_at=now,
                    updated_at=now,
                )
            )
        db.flush()

    db.commit()

    print(f"Imported/updated {len(plu_rows)} items")
    if unmatched:
        print(f"Items without 8-digit match ({len(unmatched)}):")
        for line in unmatched:
            print(f"  - {line}")


def _parse_vendor_blocks(path: str) -> List[Tuple[str, List[Tuple[str, str, str, Decimal]]]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Rates 2026"]
    vendors: List[Tuple[str, List[Tuple[str, str, str, Decimal]]]] = []
    current_vendor: Optional[str] = None
    current_items: List[Tuple[str, str, str, Decimal]] = []

    for row in ws.iter_rows(values_only=True):
        label = row[1]
        if label and row[0] is None and label not in ("Annexure", "Effective Rate", "Item Code"):
            if current_vendor and current_items:
                vendors.append((current_vendor, current_items))
            current_vendor = str(label).strip()
            current_items = []
            continue

        if not current_vendor or not row[1]:
            continue

        raw_code = row[0]
        name = str(row[1]).strip()
        if name in SKIP_VENDOR_ITEM_NAMES:
            continue
        if not raw_code:
            continue

        code = str(int(raw_code)) if isinstance(raw_code, (int, float)) else str(raw_code)
        uom = str(row[2]).strip().upper() if row[2] else ""
        rate = Decimal(str(row[3])) if row[3] is not None else Decimal("0")
        current_items.append((code, name, uom, rate))

    if current_vendor and current_items:
        vendors.append((current_vendor, current_items))
    wb.close()
    return vendors


def import_vendors(db: Session, client_data_path: str) -> None:
    blocks = _parse_vendor_blocks(client_data_path)
    now = int(time.time())
    skipped: List[str] = []

    for vendor_name, items in blocks:
        if vendor_name in SKIP_VENDOR_ITEM_NAMES:
            continue

        vendor = db.query(Vendor).filter(Vendor.name == vendor_name).first()
        if not vendor:
            vendor = Vendor(name=vendor_name, is_active=True, created_at=now)
            db.add(vendor)
            db.flush()

        for item_code_8, name, _uom, rate in items:
            item = (
                db.query(ItemMaster)
                .filter(ItemMaster.item_code == item_code_8)
                .first()
            )
            if not item:
                skipped.append(f"{vendor_name} | {item_code_8} | {name}")
                continue

            existing = (
                db.query(VendorItem)
                .filter(
                    VendorItem.vendor_id == vendor.id,
                    VendorItem.item_id == item.id,
                )
                .first()
            )
            if existing:
                existing.rate_inr = rate
            else:
                db.add(
                    VendorItem(
                        vendor_id=vendor.id,
                        item_id=item.id,
                        rate_inr=rate,
                    )
                )

    db.commit()
    print(f"Imported/updated {len(blocks)} vendor blocks")
    if skipped:
        print(f"Skipped vendor rows without item match ({len(skipped)}):")
        for line in skipped[:20]:
            print(f"  - {line}")
        if len(skipped) > 20:
            print(f"  ... +{len(skipped) - 20} more")


def main() -> None:
    parser = argparse.ArgumentParser(description="Import DigiScrapyard client master data")
    default_item = os.path.join(ROOT, DEFAULT_ITEM_CODE_XLSX)
    default_client = os.path.join(ROOT, DEFAULT_CLIENT_DATA_XLSX)
    parser.add_argument(
        "--item-code",
        default=default_item,
        help="Path to PLU item code xlsx",
    )
    parser.add_argument(
        "--client-data",
        default=default_client,
        help="Path to client 8-digit / vendor rates xlsx",
    )
    args = parser.parse_args()

    if not os.path.exists(args.item_code):
        raise SystemExit(f"Item code file not found: {args.item_code}")
    if not os.path.exists(args.client_data):
        raise SystemExit(f"Client data file not found: {args.client_data}")

    db = SessionLocal()
    try:
        import_items(db, args.item_code, args.client_data)
        import_vendors(db, args.client_data)
    finally:
        db.close()


if __name__ == "__main__":
    main()
